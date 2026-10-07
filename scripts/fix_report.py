import re

with open('scripts/generate_report.py', 'r', encoding='utf-8') as f:
    code = f.read()

# 1. Update Chapter 2 sections to be concise and well-structured
ch2_replacement = """
    SECTION_MAP = {
        '2.1 Dataset Source': [
            'The dataset used is the India YouTube Trending Videos dataset (INvideos.csv).',
            'It contains information about trending YouTube videos in India, including video details, views, likes, and comments.',
            'The project uses this dataset to perform exploratory data analysis and build an interactive dashboard.',
        ],
        '2.2 Dataset Description': [
            'The data is cleaned and preprocessed. The following columns are used:',
            'Video Information: video_id, trending_date, title, channel_title, category_id, publish_time',
            'Interaction Metrics: views, likes, dislikes, comment_count',
            'Derived Features: publish_year, publish_month, publish_day, publish_hour, engagement, engagement_rate',
            'Original unused columns (e.g., tags, thumbnail_link, description) were excluded.',
        ],
        '2.3 Features / Attributes': [
            'The attributes are categorized into core information and derived metrics.',
        ],
        '2.4 Target Variable, if Applicable': [
            'This is a descriptive data analysis project, so there is no target variable for prediction.',
            'The focus is on understanding engagement patterns and trending behavior.',
        ],
        '2.5 Dataset Statistics / Initial Exploration': [
            'Initial exploration included displaying rows (df.head, df.tail), checking data types (df.info), and identifying missing/duplicate values.',
            'The final processed dataset contains 33,089 records and 16 columns.',
        ],
    }
"""
code = re.sub(r"SECTION_MAP = \{\s*'2\.1 Dataset Source': \[.*?\n    \}", ch2_replacement.strip(), code, flags=re.DOTALL)


# 2. Add table insertion helper and fix Chapter 4 (Architecture and Flowchart)
table_helper = """
def insert_single_col_table(ref_para, items):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    table = doc.add_table(rows=len(items), cols=1)
    table.style = 'Table Grid'
    for i, item in enumerate(items):
        cell = table.cell(i, 0)
        cell.text = item
        # Center align and set font
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(11)
    # Move table to after ref_para
    ref_para._p.addnext(table._tbl)
"""

if "def insert_single_col_table" not in code:
    code = code.replace("def update_ch4(doc):", table_helper + "\n\ndef update_ch4(doc):")

ch4_arch_flowchart = """
    ARCH_FLOW = [
        'The system architecture follows a sequential data processing pipeline:',
        '[INSERT_ARCH_TABLE]'
    ]

    FLOWCHART = [
        'The flowchart for the data analysis process is as follows:',
        '[INSERT_FLOW_TABLE]'
    ]
"""

# Replace METH_FLOW, ARCH_FLOW, FLOWCHART, ALGO_SECTION definitions
ch4_repl_regex = r"METH_FLOW = \[.*?\]\s*ARCH_FLOW = \[.*?\]\s*FLOWCHART = \[.*?\]\s*ALGO_SECTION = \[.*?\]"
new_ch4 = """
    METH_FLOW = [
        'The project follows a systematic data analysis methodology.',
        'Data Loading → Data Exploration → Data Cleaning → Data Preprocessing → Feature Engineering → EDA → Visualization → Dashboard → Insights'
    ]
""" + ch4_arch_flowchart + """
    ALGO_SECTION = [
        'The primary approach used is Exploratory Data Analysis (EDA) and descriptive statistics.',
        'The analysis uses group-by operations, sorting, percentage calculations, and time-based aggregation to extract patterns.',
    ]
"""
code = re.sub(ch4_repl_regex, new_ch4.strip(), code, flags=re.DOTALL)

# Handle the table insertion in update_ch4
table_insert_logic = """
        for idx, line in enumerate(lines):
            if idx < len(paras):
                if line == '[INSERT_ARCH_TABLE]':
                    clear_para(paras[idx])
                    arch_steps = ['INvideos.csv', '↓', 'Python / Google Colab', '↓', 'Pandas Data Processing', '↓', 'Data Cleaning', '↓', 'Feature Engineering', '↓', 'Exploratory Data Analysis', '↓', 'Statistical Analysis', '↓', 'Data Visualization', '↓', 'youtube_dashboard_data.csv', '↓', 'React / Vite Dashboard', '↓', 'Insights']
                    insert_single_col_table(paras[idx], arch_steps)
                elif line == '[INSERT_FLOW_TABLE]':
                    clear_para(paras[idx])
                    flow_steps = ['START', '↓', 'Load INvideos.csv', '↓', 'Dataset Exploration', '↓', 'Missing / Duplicate Check', '↓', 'Data Cleaning & Formatting', '↓', 'Feature Engineering', '↓', 'Exploratory Data Analysis', '↓', 'Statistical Analysis', '↓', 'Data Visualization', '↓', 'Export CSV', '↓', 'React Dashboard', '↓', 'END']
                    insert_single_col_table(paras[idx], flow_steps)
                else:
                    update_para(paras[idx], line)
"""
code = re.sub(r"for idx, line in enumerate\(lines\):\s*if idx < len\(paras\):\s*update_para\(paras\[idx\], line\)", table_insert_logic.strip(), code, flags=re.DOTALL)


with open('scripts/generate_report.py', 'w', encoding='utf-8') as f:
    f.write(code)
print("Fix script applied successfully.")
