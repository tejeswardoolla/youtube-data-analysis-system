import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

doc = Document(r'C:\Users\TEJESWAR\Desktop\YouTube-Data-Analysis\DAE_Reference_Document.docx')
results = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt:
        results.append({'i': i, 'style': p.style.name, 'text': txt[:120]})

with open(r'C:\Users\TEJESWAR\Desktop\YouTube-Data-Analysis\para_dump.txt', 'w', encoding='utf-8') as f:
    for r in results:
        f.write(f"[{r['i']:03d}] [{r['style']}] {r['text']}\n")

print('Done. Total paragraphs with text:', len(results))
print('Total tables:', len(doc.tables))
for ti, table in enumerate(doc.tables):
    print(f'  Table {ti}: {len(table.rows)} rows x {len(table.columns)} cols')
    for row in table.rows[:2]:
        cells = [c.text.strip()[:40] for c in row.cells]
        print(f'    {cells}')
